import sys
from io import BytesIO
from pathlib import Path

from docx import Document
from openpyxl import Workbook

sys.path.append(str(Path(__file__).resolve().parents[1]))

from app.services.content_extraction import extract_text_from_upload


def _docx_bytes(text: str) -> bytes:
    buffer = BytesIO()
    doc = Document()
    doc.add_paragraph(text)
    doc.save(buffer)
    return buffer.getvalue()


def _xlsx_bytes() -> bytes:
    buffer = BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = "Processes"
    ws.append(["Process", "System"])
    ws.append(["Автопролонгация ИБС", "SFA"])
    wb.save(buffer)
    return buffer.getvalue()


def _text_pdf_bytes(text: str) -> bytes:
    content = f"BT /F1 12 Tf 72 72 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 144] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(content)).encode("ascii") + b" >>\nstream\n" + content + b"\nendstream",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for idx, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out.extend(f"{idx} 0 obj\n".encode("ascii"))
        out.extend(obj)
        out.extend(b"\nendobj\n")
    xref_offset = len(out)
    out.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    out.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        out.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    out.extend(
        f"trailer\n<< /Root 1 0 R /Size {len(objects) + 1} >>\nstartxref\n{xref_offset}\n%%EOF\n".encode(
            "ascii"
        )
    )
    return bytes(out)


def test_docx_extraction_returns_completed_and_chunks():
    result = extract_text_from_upload("context.docx", _docx_bytes("Описание процесса ИБС"))

    assert result["status"] == "completed"
    assert "Описание процесса ИБС" in result["text"]
    assert result["chunks"]
    assert result["chunks"][0]["metadata"]["start"] == 0


def test_txt_extraction_returns_completed_and_chunks():
    result = extract_text_from_upload("notes.txt", "Процесс использует CRM".encode("utf-8"))

    assert result["status"] == "completed"
    assert result["text"] == "Процесс использует CRM"
    assert result["chunks"][0]["text"] == "Процесс использует CRM"


def test_empty_file_returns_empty():
    result = extract_text_from_upload("empty.md", b"   ")

    assert result["status"] == "empty"
    assert result["text"] == ""
    assert result["chunks"] == []


def test_unsupported_file_returns_unsupported():
    result = extract_text_from_upload("image.png", b"not an image")

    assert result["status"] == "unsupported"
    assert "png" in result["error"]


def test_pdf_text_extraction_works_for_text_pdf():
    result = extract_text_from_upload("context.pdf", _text_pdf_bytes("Process Alpha uses CRM"))

    assert result["status"] == "completed"
    assert "Process Alpha uses CRM" in result["text"]
    assert result["chunks"]


def test_xlsx_extraction_works_for_simple_workbook():
    result = extract_text_from_upload(
        "context.xlsx",
        _xlsx_bytes(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

    assert result["status"] == "completed"
    assert "Processes" in result["text"]
    assert "Автопролонгация ИБС" in result["text"]
    assert result["chunks"]


def test_xls_returns_honest_unsupported():
    result = extract_text_from_upload("legacy.xls", b"binary")

    assert result["status"] == "unsupported"
    assert "xls" in result["error"].lower()
